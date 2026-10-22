"""
Arma las dos presentaciones de la unidad 03 sobre la plantilla de la materia.

    uv run --with python-pptx python 03-refuerzos/presentaciones/armar_presentaciones.py

Toma `utils/00_ Modelo de Presentación.pptx`: conserva su portada, su índice y
su diapositiva de referencias, y agrega las de contenido con el formato de la
presentación de búsqueda (título azul, texto gris en Roboto, notas en verde).

Las fórmulas se escriben en LaTeX y se insertan como imagen: hace falta
pdflatex y pdftocairo. Las figuras se dibujan con matplotlib, y las que
muestran resultados se sacan de las salidas guardadas en los notebooks de la
clase. Todo lo generado va a `img/`.

Además de los .pptx, deja en `vista_previa/` un PNG aproximado por diapositiva
(con DejaVu en lugar de Roboto) para revisar que nada se pise ni se salga de la
hoja sin abrir PowerPoint.
"""

import base64
import copy
import hashlib
import json
import math
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import gymnasium as gym
import matplotlib

matplotlib.use("Agg")
import matplotlib.image as mpimg
import matplotlib.patches as patches
import matplotlib.pyplot as plt
import numpy as np
from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

AQUI = Path(__file__).resolve().parent
UNIDAD = AQUI.parent
MODELO = UNIDAD / "utils" / "00_ Modelo de Presentación.pptx"
IMG = AQUI / "img"
VISTA = AQUI / "vista_previa"

# Los colores de la presentación de búsqueda
AZUL = "1C4587"
GRIS = "666666"
VERDE = "38761D"
OSCURO = "434343"
FUENTE = "Roboto"

# Zona útil de una diapositiva de contenido (pulgadas, hoja de 10 x 5.625)
X0, Y0, Y_MAX = 0.45, 1.2, 5.05
ANCHO_TOTAL = 9.1
ANCHO_CON_FIGURA = 5.5


# ---------------------------------------------------------------------------
# Fórmulas: LaTeX -> PDF -> PNG transparente
# ---------------------------------------------------------------------------

PREAMBULO = r"""
\documentclass[border=1pt]{standalone}
\usepackage{amsmath,amssymb,bm}
\usepackage[T1]{fontenc}
\usepackage{xcolor}
\begin{document}
\color[HTML]{434343}
$\displaystyle %s$
\end{document}
"""

# Un punto de LaTeX ocupa ESCALA puntos en la diapositiva: el texto de las
# fórmulas queda de unos 16 pt, un poco más grande que el cuerpo (13 pt).
ESCALA = 1.55


def formula_png(tex):
    """Compila la fórmula y devuelve (ruta del PNG, ancho, alto) en pulgadas."""
    clave = hashlib.sha1(tex.encode()).hexdigest()[:12]
    destino = IMG / "formulas" / f"{clave}.png"
    medida = destino.with_suffix(".json")
    if not destino.exists():
        destino.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "f.tex").write_text(PREAMBULO % tex)
            r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "f.tex"],
                               cwd=tmp, capture_output=True, text=True)
            if r.returncode != 0:
                raise RuntimeError(f"LaTeX falló con {tex!r}:\n{r.stdout[-1500:]}")
            info = subprocess.run(["pdfinfo", "f.pdf"], cwd=tmp, capture_output=True, text=True).stdout
            w, h = map(float, re.search(r"Page size:\s+([\d.]+) x ([\d.]+)", info).groups())
            subprocess.run(["pdftocairo", "-png", "-transp", "-r", "600", "-singlefile",
                            "f.pdf", str(destino.with_suffix(""))], cwd=tmp, check=True)
        medida.write_text(json.dumps([w, h]))
    w, h = json.loads(medida.read_text())
    return destino, w * ESCALA / 72, h * ESCALA / 72


# ---------------------------------------------------------------------------
# Figuras
# ---------------------------------------------------------------------------

def guardar(fig, nombre):
    IMG.mkdir(parents=True, exist_ok=True)
    ruta = IMG / nombre
    fig.savefig(ruta, dpi=200, bbox_inches="tight", transparent=True)
    plt.close(fig)
    return ruta


def flecha(ax, a, b, texto=None, desplazamiento=(0, 0), color="#" + OSCURO, **kw):
    ax.annotate("", xy=b, xytext=a, arrowprops=dict(arrowstyle="-|>", color=color, lw=2, **kw))
    if texto:
        ax.text((a[0] + b[0]) / 2 + desplazamiento[0], (a[1] + b[1]) / 2 + desplazamiento[1],
                texto, ha="center", va="center", fontsize=15, color="#" + OSCURO)


def caja(ax, x, y, w, h, texto, color):
    ax.add_patch(patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                                        fc=color, ec="none"))
    ax.text(x + w / 2, y + h / 2, texto, ha="center", va="center", fontsize=15, color="white",
            fontweight="bold")


def fig_agente_entorno():
    fig, ax = plt.subplots(figsize=(4.2, 3.2))
    caja(ax, 0.3, 2.2, 1.8, 0.7, "Agente", "#" + AZUL)
    caja(ax, 0.3, 0.2, 1.8, 0.7, "Entorno", "#6D9EEB")
    # acción: del agente al entorno, por la derecha
    ax.plot([2.1, 2.9, 2.9], [2.55, 2.55, 0.55], color="#" + OSCURO, lw=2)
    flecha(ax, (2.9, 0.55), (2.12, 0.55))
    ax.text(3.0, 1.55, "acción\n$A_t$", fontsize=13, color="#" + OSCURO, va="center")
    # estado y recompensa: del entorno al agente, por la izquierda
    ax.plot([0.3, -0.5, -0.5], [0.55, 0.55, 2.4], color="#" + OSCURO, lw=2)
    flecha(ax, (-0.5, 2.4), (0.28, 2.4))
    ax.plot([0.3, -0.2, -0.2], [0.75, 0.75, 2.7], color="#" + OSCURO, lw=2)
    flecha(ax, (-0.2, 2.7), (0.28, 2.7))
    ax.text(-0.6, 1.55, "estado $S_{t+1}$\nrecompensa $R_{t+1}$", fontsize=13,
            color="#" + OSCURO, va="center", ha="right")
    ax.set_xlim(-2.6, 3.8)
    ax.set_ylim(0, 3.1)
    ax.axis("off")
    return guardar(fig, "agente-entorno.png")


def dibujar_grilla_4x3(ax, politica=None, valores=None):
    terminales = {(4, 3): ("+1", "#93C47D"), (4, 2): ("−1", "#E06666")}
    for x in range(1, 5):
        for y in range(1, 4):
            color = "#CCCCCC" if (x, y) == (2, 2) else "white"
            if (x, y) in terminales:
                color = terminales[(x, y)][1]
            ax.add_patch(patches.Rectangle((x - 1, y - 1), 1, 1, fc=color, ec="#" + OSCURO, lw=1.5))
            if (x, y) in terminales:
                ax.text(x - 0.5, y - 0.5, terminales[(x, y)][0], ha="center", va="center",
                        fontsize=20, color="white", fontweight="bold")
            elif politica and (x, y) in politica:
                ax.text(x - 0.5, y - 0.5, politica[(x, y)], ha="center", va="center",
                        fontsize=26, color="#" + AZUL)
            elif valores and (x, y) in valores:
                ax.text(x - 0.5, y - 0.5, f"{valores[(x, y)]:.2f}", ha="center", va="center",
                        fontsize=11, color="#" + AZUL)
    if not politica and not valores:
        ax.text(0.5, 0.5, "S", ha="center", va="center", fontsize=20, color="#" + AZUL,
                fontweight="bold")
    for x in range(1, 5):
        ax.text(x - 0.5, -0.25, str(x), ha="center", va="center", fontsize=11, color="#" + GRIS)
    for y in range(1, 4):
        ax.text(-0.25, y - 0.5, str(y), ha="center", va="center", fontsize=11, color="#" + GRIS)
    ax.set_xlim(-0.5, 4.05)
    ax.set_ylim(-0.5, 3.05)
    ax.set_aspect("equal")
    ax.axis("off")


def fig_mundo_4x3():
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(7, 3), gridspec_kw=dict(width_ratios=[4.5, 2.6]))
    dibujar_grilla_4x3(a1)
    # el modelo de movimiento: 0.8 hacia donde se elige, 0.1 a cada costado
    a2.add_patch(patches.Circle((0, 0), 0.12, color="#" + AZUL))
    flecha(a2, (0, 0.15), (0, 1.0), color="#" + AZUL)
    flecha(a2, (-0.15, 0), (-0.9, 0), color="#" + GRIS)
    flecha(a2, (0.15, 0), (0.9, 0), color="#" + GRIS)
    a2.text(0, 1.15, "0.8", ha="center", fontsize=17, color="#" + AZUL)
    a2.text(-0.75, 0.15, "0.1", ha="center", fontsize=17, color="#" + GRIS)
    a2.text(0.75, 0.15, "0.1", ha="center", fontsize=17, color="#" + GRIS)
    a2.text(0, -0.45, "acción ↑", ha="center", fontsize=15, color="#" + OSCURO)
    a2.set_xlim(-1.15, 1.15)
    a2.set_ylim(-0.7, 1.4)
    a2.set_aspect("equal")
    a2.axis("off")
    return guardar(fig, "mundo-4x3.png")


def fig_politica_4x3():
    politica = {(1, 3): "→", (2, 3): "→", (3, 3): "→", (1, 2): "↑", (3, 2): "↑",
                (1, 1): "↑", (2, 1): "←", (3, 1): "←", (4, 1): "←"}
    fig, ax = plt.subplots(figsize=(4.2, 3.2))
    dibujar_grilla_4x3(ax, politica)
    return guardar(fig, "politica-4x3.png")


def iteracion_valores_4x3(iteraciones, costo_paso=-0.04):
    # La misma iteración de valores que el notebook de la clase 1, con γ = 1.
    movimientos = {"↑": (0, 1), "↓": (0, -1), "←": (-1, 0), "→": (1, 0)}
    costados = {"↑": ("←", "→"), "↓": ("→", "←"), "←": ("↓", "↑"), "→": ("↑", "↓")}
    terminales = {(4, 3): 1.0, (4, 2): -1.0}
    estados = [(x, y) for x in range(1, 5) for y in range(1, 4)
               if (x, y) != (2, 2) and (x, y) not in terminales]

    def mover(s, a):
        d = (s[0] + movimientos[a][0], s[1] + movimientos[a][1])
        return s if d == (2, 2) or not (1 <= d[0] <= 4 and 1 <= d[1] <= 3) else d

    def valor_accion(v, s, a):
        total = 0.0
        for prob, b in ((0.8, a), (0.1, costados[a][0]), (0.1, costados[a][1])):
            s2 = mover(s, b)
            total += prob * (terminales.get(s2, costo_paso) + v.get(s2, 0.0))
        return total

    v = {s: 0.0 for s in estados}
    for _ in range(iteraciones):
        v = {s: max(valor_accion(v, s, a) for a in movimientos) for s in estados}
    return v


def fig_iteracion_valores_4x3():
    fig, axes = plt.subplots(1, 3, figsize=(10, 2.9))
    for ax, k in zip(axes, (1, 3, 30)):
        dibujar_grilla_4x3(ax, valores=iteracion_valores_4x3(k))
        ax.set_title(f"iteración {k}" if k < 30 else "convergido", fontsize=14, color="#" + OSCURO)
    return guardar(fig, "valores-4x3.png")


def figuras_del_resumen(*nombres):
    """Las figuras de `resumen/figuras.py`, como PNG en img/. Devuelve sus rutas."""
    import importlib.util
    import sys

    rutas = {}

    def a_png(fig, nombre):
        rutas[nombre] = IMG / f"{nombre}.png"
        fig.savefig(rutas[nombre], dpi=200, bbox_inches="tight", transparent=True)
        plt.close(fig)

    # resumen/figuras.py cambia los rcParams al importarse: se aíslan para que no
    # alteren el resto de las figuras de las presentaciones.
    with plt.rc_context():
        spec = importlib.util.spec_from_file_location("figuras_resumen", UNIDAD / "resumen" / "figuras.py")
        modulo = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(modulo)
        matplotlib.use("Agg")
        modulo.guardar = a_png
        IMG.mkdir(parents=True, exist_ok=True)
        for nombre in nombres:
            getattr(modulo, "fig_" + nombre.replace("-", "_"))()
        sys.modules.pop("figuras_resumen", None)
    return rutas


def entrenar_cliff(metodo, episodios=500, alpha=0.5, epsilon=0.1, semilla=0):
    # La misma función que el notebook de la clase 2.
    rng = np.random.default_rng(semilla)
    env = gym.make("CliffWalking-v1")
    Q = np.zeros((48, 4))

    def elegir(s):
        if rng.random() < epsilon:
            return int(rng.integers(4))
        return int(rng.choice(np.flatnonzero(Q[s] == Q[s].max())))

    for ep in range(episodios):
        s, _ = env.reset(seed=semilla * 1000 + ep)
        a = elegir(s)
        fin = False
        while not fin:
            s2, r, terminado, truncado, _ = env.step(a)
            a2 = elegir(s2)
            siguiente = Q[s2].max() if metodo == "q-learning" else Q[s2, a2]
            Q[s, a] += alpha * (r + (0 if terminado else siguiente) - Q[s, a])
            s, a, fin = s2, a2, terminado or truncado
    return Q


def camino_greedy(Q):
    env = gym.make("CliffWalking-v1")
    s, _ = env.reset()
    camino = [s]
    for _ in range(100):
        s, _, terminado, truncado, _ = env.step(int(np.argmax(Q[s])))
        camino.append(s)
        if terminado or truncado:
            break
    return camino


def fig_cliff():
    fig, ax = plt.subplots(figsize=(7.5, 2.9))
    for fila in range(4):
        for col in range(12):
            y = 3 - fila
            color = "white"
            if fila == 3 and 0 < col < 11:
                color = "#999999"
            ax.add_patch(patches.Rectangle((col, y), 1, 1, fc=color, ec="#CCCCCC", lw=1))
    ax.text(6, 0.5, "acantilado (−100)", ha="center", va="center", fontsize=13, color="white")
    ax.text(0.5, 0.5, "S", ha="center", va="center", fontsize=15, color="#" + OSCURO, fontweight="bold")
    ax.text(11.5, 0.5, "G", ha="center", va="center", fontsize=15, color="#" + OSCURO, fontweight="bold")
    for metodo, color, dy in (("q-learning", "#CC4125", -0.08), ("sarsa", "#" + AZUL, 0.08)):
        camino = camino_greedy(entrenar_cliff(metodo))
        xs = [s % 12 + 0.5 for s in camino]
        ys = [3 - s // 12 + 0.5 + dy for s in camino]
        ax.plot(xs, ys, color=color, lw=3, label=metodo.capitalize().replace("Sarsa", "SARSA"))
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.02), ncol=2, frameon=False, fontsize=13)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4)
    ax.set_aspect("equal")
    ax.axis("off")
    return guardar(fig, "cliff-caminos.png")


def fig_cartpole():
    env = gym.make("CartPole-v1", render_mode="rgb_array")
    env.reset(seed=3)
    for _ in range(6):
        env.step(1)
    cuadro = env.render()[90:330, 60:540]
    fig, ax = plt.subplots(figsize=(4.8, 2.4))
    ax.imshow(cuadro)
    ax.axis("off")
    return guardar(fig, "cartpole.png")


def fig_red_q():
    fig, ax = plt.subplots(figsize=(4.6, 3.4))
    capas = [4, 6, 6, 2]
    xs = [0, 1.4, 2.8, 4.2]
    entradas = [r"$\dot{\theta}$", r"$\theta$", r"$\dot{x}$", "$x$"]   # de abajo hacia arriba
    salidas = [r"$\hat{q}(s, \leftarrow)$", r"$\hat{q}(s, \rightarrow)$"]
    posiciones = []
    for n, x in zip(capas, xs):
        ys = np.linspace(-(n - 1) / 2, (n - 1) / 2, n) * 0.55
        posiciones.append([(x, y) for y in ys])
    for capa_a, capa_b in zip(posiciones, posiciones[1:]):
        for a in capa_a:
            for b in capa_b:
                ax.plot([a[0], b[0]], [a[1], b[1]], color="#CCCCCC", lw=0.6, zorder=1)
    colores = ["#6D9EEB", "#BBBBBB", "#BBBBBB", "#" + AZUL]
    for capa, color in zip(posiciones, colores):
        for x, y in capa:
            ax.add_patch(patches.Circle((x, y), 0.16, color=color, zorder=2))
    for (x, y), t in zip(posiciones[0], entradas):
        ax.text(x - 0.3, y, t, ha="right", va="center", fontsize=14, color="#" + OSCURO)
    for (x, y), t in zip(posiciones[-1], salidas):
        ax.text(x + 0.3, y, t, ha="left", va="center", fontsize=14, color="#" + OSCURO)
    ax.text(0, -1.45, "estado $s$", ha="center", fontsize=12, color="#" + GRIS)
    ax.text(2.1, -1.95, "capas ocultas", ha="center", fontsize=12, color="#" + GRIS)
    ax.text(4.2, -1.45, "una salida\npor acción", ha="center", va="top", fontsize=12, color="#" + GRIS)
    ax.set_xlim(-1.2, 6.3)
    ax.set_ylim(-2.2, 1.8)
    ax.set_aspect("equal")
    ax.axis("off")
    return guardar(fig, "red-q.png")


def fig_memoria():
    fig, ax = plt.subplots(figsize=(4.6, 3.3))
    caja(ax, -0.4, 2.6, 1.9, 0.6, "Entorno", "#6D9EEB")
    caja(ax, 3.1, 2.6, 1.5, 0.6, "Red", "#" + AZUL)
    # la memoria: una pila de transiciones
    rng = np.random.default_rng(1)
    elegidas = set(rng.choice(8, 3, replace=False))
    for i in range(8):
        y = 0.1 + i * 0.27
        color = "#F6B26B" if i in elegidas else "#EEEEEE"
        ax.add_patch(patches.Rectangle((1.35, y), 1.9, 0.22, fc=color, ec="#BBBBBB"))
        ax.text(2.3, y + 0.11, f"(s, a, r, s′)", ha="center", va="center", fontsize=8, color="#" + OSCURO)
    ax.text(2.3, -0.25, "memoria $D$", ha="center", fontsize=12, color="#" + GRIS)
    flecha(ax, (0.75, 2.58), (1.3, 2.1), "guardar", (-0.75, -0.3))
    flecha(ax, (3.3, 1.2), (3.85, 2.55), "minibatch\nal azar", (0.95, -0.45))
    ax.set_xlim(-0.8, 5.6)
    ax.set_ylim(-0.4, 3.3)
    ax.axis("off")
    return guardar(fig, "memoria.png")


def png_de_notebook(notebook, marca, indice=-1, nombre=None):
    """El indice-ésimo PNG de la primera celda de código que contiene `marca`."""
    nb = json.loads((UNIDAD / notebook).read_text())
    for celda in nb["cells"]:
        if celda["cell_type"] == "code" and marca in "".join(celda["source"]):
            pngs = [o["data"]["image/png"] for o in celda.get("outputs", [])
                    if "image/png" in o.get("data", {})]
            if not pngs:
                raise RuntimeError(f"la celda con {marca!r} de {notebook} no tiene salidas: ¿se ejecutó?")
            ruta = IMG / nombre
            IMG.mkdir(parents=True, exist_ok=True)
            ruta.write_bytes(base64.b64decode(pngs[indice]))
            return ruta
    raise RuntimeError(f"no encuentro {marca!r} en {notebook}")


# ---------------------------------------------------------------------------
# Elementos de una diapositiva y su disposición
# ---------------------------------------------------------------------------

@dataclass
class Texto:
    parrafos: list
    tam: float = 13
    color: str = GRIS
    vinetas: bool = False


@dataclass
class Nota:
    texto: str
    color: str = VERDE
    tam: float = 12


@dataclass
class Formula:
    tex: str
    escala: float = 1.0


@dataclass
class Codigo:
    texto: str
    tam: float = 10


@dataclass
class Diapositiva:
    titulo: str
    columna: list
    figura: Path | None = None
    ancho_figura: float = 3.4
    figura_abajo: bool = False     # la figura va debajo del texto, a lo ancho
    ubicados: list = field(default_factory=list)


# Para saber cuántas líneas ocupa un párrafo se mide el texto con la misma
# fuente que usa la presentación. Si Roboto no está instalada, se estima.
ROBOTO = Path("/usr/share/fonts/truetype/roboto/unhinted/RobotoTTF/Roboto-Regular.ttf")
_FUENTE_PIL = ImageFont.truetype(str(ROBOTO), 100) if ROBOTO.exists() else None
INTERLINEADO = 1.2      # alto de línea, en tamaños de fuente
ESPACIO_PARRAFO = 4     # puntos después de cada párrafo


def ancho_en_pulgadas(texto, tam):
    if _FUENTE_PIL is None:
        return len(texto) * tam / 72 * 0.5
    return _FUENTE_PIL.getlength(texto) / 100 * tam / 72


def partir(texto, ancho, tam):
    """Corta un párrafo en líneas que entran en `ancho` pulgadas."""
    lineas, actual = [], ""
    for palabra in texto.replace("**", "").split():
        prueba = f"{actual} {palabra}".strip()
        if actual and ancho_en_pulgadas(prueba, tam) > ancho:
            lineas.append(actual)
            actual = palabra
        else:
            actual = prueba
    return lineas + [actual]


def alto_texto(parrafos, ancho, tam, vinetas=False):
    util = ancho - 0.1 - (0.25 if vinetas else 0)
    n = sum(len(partir(p, util, tam)) for p in parrafos)
    return n * tam / 72 * INTERLINEADO + len(parrafos) * ESPACIO_PARRAFO / 72 + 0.08


def disponer(d):
    """Calcula la posición de cada elemento. Devuelve la altura en que termina."""
    ancho = ANCHO_TOTAL if (d.figura is None or d.figura_abajo) else ANCHO_CON_FIGURA
    y = Y0
    d.ubicados = []
    for el in d.columna:
        if isinstance(el, Texto):
            h = alto_texto(el.parrafos, ancho, el.tam, el.vinetas)
            d.ubicados.append((el, X0, y, ancho, h))
        elif isinstance(el, Nota):
            h = alto_texto([el.texto], ancho, el.tam)
            d.ubicados.append((el, X0, y, ancho, h))
        elif isinstance(el, Codigo):
            n = len(el.texto.strip("\n").splitlines())
            h = n * (el.tam * INTERLINEADO + 2) / 72 + 0.1
            d.ubicados.append((el, X0, y, ancho, h))
        elif isinstance(el, Formula):
            png, w, h = formula_png(el.tex)
            w, h = w * el.escala, h * el.escala
            if w > ancho - 0.2:
                h *= (ancho - 0.2) / w
                w = ancho - 0.2
            d.ubicados.append((el, X0 + (ancho - w) / 2, y + 0.04, w, h))
            h += 0.08
        y += h + 0.1
    if d.figura is not None:
        ancho_png, alto_png = plt.imread(d.figura).shape[1::-1]
        if d.figura_abajo:
            disponible = Y_MAX - y
            h = min(disponible, d.ancho_figura * alto_png / ancho_png)
            w = h * ancho_png / alto_png
            d.ubicados.append((d.figura, X0 + (ANCHO_TOTAL - w) / 2, y, w, h))
            y += h
        else:
            w = d.ancho_figura
            h = w * alto_png / ancho_png
            if h > 3.7:
                w, h = w * 3.7 / h, 3.7
            x = X0 + ANCHO_CON_FIGURA + 0.2 + (ANCHO_TOTAL - ANCHO_CON_FIGURA - 0.2 - w) / 2
            d.ubicados.append((d.figura, x, Y0 + (Y_MAX - Y0 - h) / 2, w, h))
    return y


# ---------------------------------------------------------------------------
# Escritura en el .pptx
# ---------------------------------------------------------------------------

def agregar_runs(parrafo, texto, tam, color, negrita=False, fuente=FUENTE):
    # **texto** va en negrita, como en el resto del material
    for i, pedazo in enumerate(re.split(r"\*\*", texto)):
        if not pedazo:
            continue
        run = parrafo.add_run()
        run.text = pedazo
        run.font.name = fuente
        run.font.size = Pt(tam)
        run.font.bold = negrita or (i % 2 == 1)
        run.font.color.rgb = RGBColor.from_string(color)


def con_vineta(parrafo):
    pPr = parrafo._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.25)))
    pPr.set("indent", str(-Inches(0.18)))
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "•")


def caja_de_texto(slide, x, y, w, h, parrafos, tam, color, vinetas=False, fuente=FUENTE,
                  espacio=ESPACIO_PARRAFO):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, texto in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(espacio)
        if vinetas:
            con_vineta(p)
        agregar_runs(p, texto, tam, color, fuente=fuente)
    return tb


def escribir(slide, d):
    slide.shapes.title.text = d.titulo
    for el, x, y, w, h in d.ubicados:
        if isinstance(el, Texto):
            caja_de_texto(slide, x, y, w, h, el.parrafos, el.tam, el.color, el.vinetas)
        elif isinstance(el, Nota):
            caja_de_texto(slide, x, y, w, h, [el.texto], el.tam, el.color)
        elif isinstance(el, Codigo):
            caja_de_texto(slide, x, y, w, h, el.texto.strip("\n").splitlines(), el.tam,
                          OSCURO, fuente="Roboto Mono", espacio=2)
        elif isinstance(el, Formula):
            slide.shapes.add_picture(str(formula_png(el.tex)[0]), Inches(x), Inches(y),
                                     Inches(w), Inches(h))
        else:  # figura
            slide.shapes.add_picture(str(el), Inches(x), Inches(y), Inches(w), Inches(h))


def reemplazar_texto(forma, parrafos, tam=None, niveles=None):
    """Cambia el texto de una forma conservando el formato de su primer run."""
    tf = forma.text_frame
    modelo_p = copy.deepcopy(tf.paragraphs[0]._p)
    for r in modelo_p.findall(qn("a:r"))[1:]:
        modelo_p.remove(r)
    txBody = tf._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for i, texto in enumerate(parrafos):
        p = copy.deepcopy(modelo_p)
        txBody.append(p)
        par = tf.paragraphs[-1]
        if par.runs:
            par.runs[0].text = texto
            if tam:
                par.runs[0].font.size = Pt(tam)
        else:
            agregar_runs(par, texto, tam or 13, GRIS)
        if niveles:
            par.level = niveles[i]


def armar(nombre, titulo, indice, diapositivas, referencias):
    prs = Presentation(MODELO)
    portada, temario, vacia, refs = list(prs.slides)
    layout_contenido = vacia.slide_layout

    # Portada: el título de la clase, un poco más ancho que en el modelo
    for forma in portada.shapes:
        if forma.has_text_frame and forma.text_frame.text.strip() == "Título de la clase":
            reemplazar_texto(forma, [titulo])
            forma.width = Inches(4.3)
            forma.height = Inches(0.85)
            forma.top = Inches(2.15)

    # Índice: temas en el nivel 0, subtemas en el 1
    cuerpo = [ph for ph in temario.placeholders if ph.placeholder_format.idx == 1][0]
    reemplazar_texto(cuerpo, [t for t, _ in indice], tam=13, niveles=[n for _, n in indice])

    # Referencias, en formato APA
    for forma in refs.shapes:
        if forma.has_text_frame and "Formato APA" in forma.text_frame.text:
            reemplazar_texto(forma, referencias, tam=11)
            forma.height = Inches(3.5)

    # La diapositiva vacía del modelo se usa para la primera de contenido; las
    # demás se agregan con su mismo diseño. (Borrarla y agregar otras hace que
    # python-pptx repita nombres de partes dentro del .pptx.)
    VISTA.mkdir(parents=True, exist_ok=True)
    for i, d in enumerate(diapositivas, start=3):
        fin = disponer(d)
        if fin > Y_MAX + 0.01:
            print(f"  ATENCIÓN: '{d.titulo}' llega hasta {fin:.2f} in (máximo {Y_MAX})")
        slide = vacia if i == 3 else prs.slides.add_slide(layout_contenido)
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx == 1:
                ph._element.getparent().remove(ph._element)
        escribir(slide, d)
        vista_previa(d, VISTA / f"{nombre[:4]}-{i:02d}.png")

    # Las referencias van al final
    lista = prs.slides._sldIdLst
    ref = list(lista)[3]
    lista.remove(ref)
    lista.append(ref)

    destino = AQUI / f"{nombre}.pptx"
    prs.save(destino)
    print(f"{destino.name}: {len(prs.slides)} diapositivas")


# ---------------------------------------------------------------------------
# Vista previa aproximada
# ---------------------------------------------------------------------------

FONDO = None


def vista_previa(d, ruta):
    global FONDO
    if FONDO is None:
        import zipfile
        with zipfile.ZipFile(MODELO) as z:
            with z.open("ppt/media/image5.png") as f:
                FONDO = mpimg.imread(f)
    fig = plt.figure(figsize=(10, 5.625), dpi=80)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.imshow(FONDO, extent=[0, 10, 5.625, 0])
    ax.set_xlim(0, 10)
    ax.set_ylim(5.625, 0)
    ax.axis("off")
    ax.add_patch(patches.Rectangle((9.36, 5.0), 0.5, 0.56, fc="#DDDDDD"))   # logo
    ax.text(0.39, 0.82, d.titulo, fontsize=25.2, color="#" + AZUL, va="center", family="Roboto")
    for el, x, y, w, h in d.ubicados:
        ax.add_patch(patches.Rectangle((x, y), w, h, fill=False, ec="#EEAAAA", lw=0.5, ls="--"))
        if isinstance(el, (Texto, Nota, Codigo)):
            if isinstance(el, Texto):
                parrafos, tam, color = el.parrafos, el.tam, el.color
            elif isinstance(el, Nota):
                parrafos, tam, color = [el.texto], el.tam, el.color
            else:
                parrafos, tam, color = el.texto.strip("\n").splitlines(), el.tam, OSCURO
            vinetas = getattr(el, "vinetas", False)
            util = w - 0.1 - (0.25 if vinetas else 0)
            codigo = isinstance(el, Codigo)
            yy = y + 0.04
            for p in parrafos:
                partes = [p] if codigo else partir(p, util, tam)
                for j, linea in enumerate(partes):
                    if vinetas and j == 0:
                        ax.text(x + 0.12, yy, "•", fontsize=tam, color="#" + color, va="top")
                    ax.text(x + 0.05 + (0.25 if vinetas else 0), yy, linea, fontsize=tam,
                            color="#" + color, va="top",
                            family="DejaVu Sans Mono" if codigo else "Roboto")
                    yy += tam / 72 * INTERLINEADO
                yy += (2 if codigo else ESPACIO_PARRAFO) / 72
        else:
            png = formula_png(el.tex)[0] if isinstance(el, Formula) else el
            ax.imshow(mpimg.imread(png), extent=[x, x + w, y + h, y])
    fig.savefig(ruta)
    plt.close(fig)


# ---------------------------------------------------------------------------
# Contenido
# ---------------------------------------------------------------------------

REF_RN = ("Russell, S., & Norvig, P. (2021). Artificial Intelligence: A Modern Approach "
          "(4th ed., Global ed.). Pearson.")
REF_SB = ("Sutton, R. S., & Barto, A. G. (2018). Reinforcement Learning: An Introduction "
          "(2nd ed.). MIT Press. http://incompleteideas.net/book/the-book-2nd.html")


def clase1():
    agente = fig_agente_entorno()
    mundo = fig_mundo_4x3()
    politica = fig_politica_4x3()
    valores = fig_iteracion_valores_4x3()
    resumen = figuras_del_resumen("tragamonedas", "mundo-4x3", "descuento", "respaldo",
                                  "volver-a-casa")
    resumen.pop("mundo4x3-utilidades").unlink()     # fig_mundo_4x3 dibuja las dos
    apostador = png_de_notebook("clase1-decisiones-secuenciales.ipynb", "HEAD_PROB = 0.4", -1,
                                "apostador.png")

    d = [
        Diapositiva("Aprendizaje por refuerzo", [
            Texto(["Un **agente** aprende qué hacer —una **política**— interactuando con un "
                   "**entorno**, para maximizar la **recompensa** que recibe a largo plazo."]),
            Texto(["En cada paso el agente observa el estado y elige una acción.",
                   "El entorno responde con una recompensa y un nuevo estado.",
                   "Nadie le dice al agente cuál era la acción correcta: solo recibe recompensas, "
                   "y muchas veces llegan tarde."], vinetas=True),
            Nota("El objetivo es encontrar la política que maximiza la recompensa acumulada."),
        ], figura=agente, ancho_figura=3.3),

        Diapositiva("Notación", [
            Texto(["Usamos la notación de Sutton y Barto. Russell y Norvig llaman **utilidad** al "
                   "valor de un estado y escriben el modelo por separado:"]),
            Texto(["Uπ(s) y U(s) en Russell y Norvig son vπ(s) y v*(s).",
                   "Q(s, a) en Russell y Norvig es q*(s, a).",
                   "P(s′ | s, a) y R(s, a, s′) en Russell y Norvig son la dinámica p(s′, r | s, a).",
                   "Las mayúsculas V y Q son estimaciones que el agente va aprendiendo."],
                  vinetas=True),
            Nota("Excepción: Sutton y Barto no tratan la teoría de la utilidad. En la sección de "
                 "decisiones bajo incertidumbre, U es la utilidad de un resultado, como en "
                 "Russell y Norvig."),
        ]),

        Diapositiva("Decisiones bajo incertidumbre", [
            Texto(["En búsqueda, cada acción tenía un resultado seguro. Ahora la misma acción puede "
                   "llevar a resultados distintos, cada uno con cierta probabilidad.",
                   "Para decidir hace falta una **utilidad** U(s), que mide qué tan deseable es cada "
                   "resultado, y la probabilidad p(s′ | a) de cada resultado."]),
            Texto(["Utilidad esperada de una acción:"]),
            Formula(r"EU(a) = \sum_{s'} p(s' \mid a)\, U(s')"),
            Texto(["Principio de **máxima utilidad esperada** (MEU): un agente racional elige"]),
            Formula(r"a^* = \arg\max_a EU(a)"),
        ]),

        Diapositiva("Ejemplo: dos apuestas", [
            Texto(["**A:** recibir 50 pesos seguros.",
                   "**B:** tirar una moneda; si sale cara se reciben 120 pesos, si sale ceca, nada."],
                  vinetas=True),
            Formula(r"EU(A) = 50 \qquad EU(B) = 0.5 \cdot 120 + 0.5 \cdot 0 = 60"),
            Texto(["Con el dinero como utilidad, conviene B. Eso no quiere decir que B gane siempre: "
                   "quiere decir que, jugando muchas veces, **en promedio** rinde más."]),
            Nota("La utilidad no tiene por qué ser el dinero: expresa las preferencias del agente. "
                 "Quien necesita sí o sí 50 pesos hoy puede preferir A sin ser irracional."),
        ]),

        Diapositiva("Tragamonedas: el valor de una acción", [
            Texto(["Hay k máquinas. Cada una paga una recompensa aleatoria, con una media propia que "
                   "el agente **no conoce**. Sin las probabilidades no se puede calcular la utilidad "
                   "esperada: hay que estimarla jugando."]),
            Texto(["El valor estimado de una acción es el promedio de las recompensas que dio:"]),
            Formula(r"Q_t(a) = \frac{\text{suma de recompensas obtenidas con } a}"
                    r"{\text{cantidad de veces que se eligió } a}"),
            Texto(["Se puede actualizar en forma incremental, sin guardar todas las recompensas:"]),
            Formula(r"Q_{n+1} = Q_n + \frac{1}{n} \left[ R_n - Q_n \right]"),
            Nota("Es la forma que se repite en toda la unidad: nueva ← vieja + paso · [objetivo − "
                 "vieja]. TD, Q-learning y DQN cambian solo el objetivo."),
        ]),

        Diapositiva("Exploración y explotación", [
            Texto(["**Explotar:** elegir la acción que hoy parece mejor (la acción greedy).",
                   "**Explorar:** probar otras acciones para estimar mejor su valor."], vinetas=True),
            Formula(r"A_t = \arg\max_a Q_t(a)"),
            Texto(["La política greedy puede quedarse para siempre en una máquina que no es la mejor, "
                   "porque nunca vuelve a probar las otras. La política **ε-greedy** explora una "
                   "fracción ε de las veces:"]),
            Formula(r"A_t = \begin{cases} \arg\max_a Q_t(a) & \text{con probabilidad } 1 - \varepsilon \\"
                    r" \text{una acción al azar} & \text{con probabilidad } \varepsilon \end{cases}"),
        ]),

        Diapositiva("UCB: explorar lo que menos se conoce", [
            Texto(["UCB (Upper Confidence Bound) no explora al azar: suma a cada valor estimado "
                   "un término que premia a las acciones poco probadas."]),
            Formula(r"A_t = \arg\max_a \left[ Q_t(a) + c \sqrt{\frac{\ln t}{N_t(a)}} \right]"),
            Texto(["Qₜ(a) premia a las acciones que parecen buenas.",
                   "La raíz crece con el tiempo t y baja con las veces Nₜ(a) que se eligió la acción: "
                   "premia a las acciones poco probadas.",
                   "c regula cuánto se explora."], vinetas=True),
            Nota("Una acción que casi no se probó tiene una estimación incierta: puede ser mejor de "
                 "lo que parece."),
        ]),

        Diapositiva("Comparación con 10 máquinas", [
            Texto(["Promedio de 2000 problemas de 10 máquinas. Greedy se estanca y elige la mejor "
                   "máquina solo un 35 % de las veces; ε-greedy y UCB siguen mejorando."]),
            Nota("La acción greedy suele ser la que tuvo suerte en las primeras jugadas: elegir el "
                 "máximo favorece a las estimaciones que salieron altas por azar."),
        ], figura=resumen["tragamonedas"], ancho_figura=8.0, figura_abajo=True),

        Diapositiva("Decisiones secuenciales: el mundo 4×3", [
            Texto(["Ahora cada acción lleva a un nuevo estado, y desde ahí hay que volver a decidir.",
                   "El agente empieza en S y el juego termina en +1 o en −1. Las acciones **no son "
                   "confiables**: van hacia donde se eligió con probabilidad 0.8, y hacia cada costado "
                   "con 0.1."]),
            Texto(["El plan ↑ ↑ → → → solo sale bien si las cinco acciones salen bien:"]),
            Formula(r"0.8^5 \approx 0.33"),
            Nota("Un plan fijo no alcanza: cuando una acción falla, el agente queda en un estado que "
                 "el plan no previó."),
        ], figura=mundo, ancho_figura=3.4),

        Diapositiva("Políticas", [
            Texto(["Una **política** dice qué hacer en **cada estado** al que se pueda llegar. Si una "
                   "acción sale mal, el agente consulta la política en el estado donde quedó."]),
            Texto(["Determinista: una acción por estado."], vinetas=True),
            Formula(r"a = \pi(s)"),
            Texto(["Estocástica: una probabilidad para cada acción."], vinetas=True),
            Formula(r"\pi(a \mid s) = P(A_t = a \mid S_t = s)"),
            Nota("La figura es la política óptima del mundo 4×3 cuando cada paso cuesta 0.04. Desde "
                 "(3,1) da toda la vuelta para no arriesgarse a caer en −1."),
        ], figura=politica, ancho_figura=3.2),

        Diapositiva("La política depende de la recompensa", [
            Texto(["Políticas óptimas del mundo 4×3 para cuatro valores de la recompensa por paso r. "
                   "Cuanto más cuesta vivir, más riesgo acepta el agente para terminar rápido."]),
            Texto(["r = −2: desde (3,2) se tira al −1, que cuesta menos que seguir caminando.",
                   "r = −0.6: toma el camino corto, aunque pase al lado del −1.",
                   "r = −0.04: da la vuelta segura.",
                   "r = −0.01: evita el −1 a toda costa; desde (4,1) choca contra la pared."],
                  vinetas=True),
            Nota("Con r > 0 vivir conviene: el agente evita las dos salidas para siempre."),
        ], figura=resumen["mundo4x3-politicas"], ancho_figura=8.0, figura_abajo=True),

        Diapositiva("Recompensas y retorno", [
            Texto(["Después de cada acción el agente recibe una recompensa. La utilidad de una "
                   "secuencia es el **retorno**: la suma de las recompensas futuras, descontadas."]),
            Formula(r"G_t = R_{t+1} + \gamma R_{t+2} + \gamma^2 R_{t+3} + \cdots "
                    r"= \sum_{k=0}^{\infty} \gamma^k R_{t+k+1} = R_{t+1} + \gamma\, G_{t+1}"),
            Texto(["El **factor de descuento** γ ∈ [0, 1] indica cuánto valen las recompensas "
                   "futuras frente a las inmediatas.",
                   "**Tareas episódicas** (terminan siempre): se puede usar γ = 1.",
                   "**Tareas continuas** (no terminan): hace falta γ < 1 para que la suma sea finita. "
                   "Si todas las recompensas valen R, el retorno es R / (1 − γ)."],
                  vinetas=True),
            Nota("Con γ cerca de 0 el agente es miope; con γ cerca de 1, previsor. Mira "
                 "unos 1 / (1 − γ) pasos: con γ = 0.9, unos 10."),
        ], figura=resumen["descuento"], ancho_figura=3.6),

        Diapositiva("Ejemplo: el retorno en el mundo 4×3", [
            Texto(["Cada paso da −0.04 y entrar en una salida da +1 o −1. Tres episodios desde (1,1):"]),
            Formula(r"\begin{array}{lccc}"
                    r" & \text{pasos} & \gamma = 1 & \gamma = 0.9 \\[2pt] \hline"
                    r"\text{camino corto a } {+1} & 5 & 4(-0.04) + 1 = 0.84 & 0.52 \\"
                    r"\text{camino largo a } {+1} & 7 & 6(-0.04) + 1 = 0.76 & 0.34 \\"
                    r"\text{cae en } {-1} & 5 & 4(-0.04) - 1 = -1.16 & -0.79"
                    r"\end{array}"),
            Texto(["Con γ < 1 el camino largo pierde más: tardar cuesta."]),
            Nota("Como las acciones fallan al azar, la misma política da episodios distintos y el "
                 "retorno es aleatorio. Por eso una política se evalúa por su retorno **esperado**."),
        ]),

        Diapositiva("Procesos de decisión de Markov", [
            Texto(["Un **MDP** queda definido por los estados, las acciones, el factor de descuento γ "
                   "y la **dinámica**: la probabilidad de cada próximo estado y recompensa."]),
            Formula(r"p(s', r \mid s, a) = \Pr\{S_{t+1} = s',\ R_{t+1} = r \mid S_t = s,\ A_t = a\}"),
            Texto(["**Propiedad de Markov:** el próximo estado y la recompensa dependen solo del "
                   "estado y la acción actuales, no de toda la historia."]),
            Nota("Las tragamonedas son un MDP con un único estado."),
        ]),

        Diapositiva("Funciones de valor", [
            Texto(["**Valor de estado:** el retorno esperado desde s, siguiendo la política π."]),
            Formula(r"v_\pi(s) = \mathbb{E}_\pi\left[ G_t \mid S_t = s \right]"),
            Texto(["**Valor de acción:** el retorno esperado si en s se toma a, y después se sigue π."]),
            Formula(r"q_\pi(s, a) = \mathbb{E}_\pi\left[ G_t \mid S_t = s, A_t = a \right]"),
            Texto(["Se relacionan en los dos sentidos:"]),
            Formula(r"v_\pi(s) = \sum_a \pi(a \mid s)\, q_\pi(s, a) \qquad "
                    r"q_\pi(s, a) = \sum_{s', r} p(s', r \mid s, a) \left[ r + \gamma\, v_\pi(s') \right]"),
            Nota("La primera promedia sobre lo que elige la política; la segunda, sobre lo que "
                 "responde el entorno."),
        ]),

        Diapositiva("Ecuación de Bellman", [
            Texto(["Reemplazando q por su expresión, el valor de un estado queda en función de los "
                   "valores de los estados que le siguen:"]),
            Formula(r"v_\pi(s) = \sum_a \pi(a \mid s) \sum_{s', r} p(s', r \mid s, a)"
                    r" \left[ r + \gamma\, v_\pi(s') \right]"),
            Texto(["π(a | s): cuánto elige la política cada acción.",
                   "p(s′, r | s, a): cómo responde el entorno.",
                   "r + γ v(s′): la recompensa inmediata más el valor descontado de donde se llega."],
                  vinetas=True),
            Nota("Es un sistema de ecuaciones **lineales**, una por estado: v = r + γPv, o sea "
                 "(I − γP) v = r. Para una política fija se puede resolver directo."),
        ]),

        Diapositiva("Ecuación de optimalidad de Bellman", [
            Texto(["La **política óptima** π* obtiene el mayor valor en todos los estados. Sus valores "
                   "cumplen la misma ecuación, pero tomando la mejor acción en lugar de promediar:"]),
            Formula(r"v_*(s) = \max_a \sum_{s', r} p(s', r \mid s, a) \left[ r + \gamma\, v_*(s') \right]"),
            Formula(r"q_*(s, a) = \sum_{s', r} p(s', r \mid s, a) \left[ r + \gamma \max_{a'} q_*(s', a') \right]"),
            Texto(["Conociendo q*, la política óptima sale directo:"]),
            Formula(r"\pi_*(s) = \arg\max_a q_*(s, a)"),
            Nota("Por el máximo, el sistema deja de ser lineal y no se resuelve directo: hay que "
                 "iterar. Con q* se elige la acción sin conocer el modelo del mundo."),
        ]),

        Diapositiva("Diagramas de respaldo", [
            Texto(["Las cuatro ecuaciones se pueden dibujar. Los círculos vacíos son estados y los "
                   "llenos, pares estado–acción. El valor **sube** desde las hojas hasta la raíz."]),
            Texto(["En vπ y qπ se **promedia**: según π al elegir la acción y según p al ver qué "
                   "responde el entorno.",
                   "En v* y q* la elección de la acción toma el **máximo**."], vinetas=True),
        ], figura=resumen["diagramas-respaldo"], ancho_figura=8.0, figura_abajo=True),

        Diapositiva("Iteración de valores", [
            Texto(["Calcula v* cuando se conocen las probabilidades de transición y las recompensas."]),
            Texto(["Inicializar v₀(s) = 0 para todos los estados.",
                   "Para cada estado, aplicar la actualización de Bellman:",
                   ], vinetas=True),
            Formula(r"v_{k+1}(s) = \max_a \sum_{s', r} p(s', r \mid s, a) \left[ r + \gamma\, v_k(s') \right]"),
            Texto(["Repetir hasta que los valores casi no cambien.",
                   "La política óptima es la acción que alcanza el máximo en cada estado."], vinetas=True),
            Nota("Necesita el modelo completo y recorrer todos los estados en cada iteración. "
                 "La iteración de políticas alterna dos pasos: evaluar la política y mejorarla."),
        ]),

        Diapositiva("Ejemplo: iteración de valores en el mundo 4×3", [
            Texto(["Con costo por paso 0.04 y γ = 1. Los valores arrancan en 0 y el +1 se va "
                   "propagando desde la salida, una casilla por iteración. La política que elige "
                   "en cada estado la acción de mayor valor es la que vimos antes."]),
            Nota("En (3,1), ir a la izquierda vale 0.65 y subir, 0.63: dar la vuelta conviene. "
                 "La política greedy suele estabilizarse mucho antes que los valores."),
        ], figura=valores, ancho_figura=9.0, figura_abajo=True),

        Diapositiva("Ejemplo: el problema del apostador", [
            Texto(["Un jugador apuesta a cara o ceca. Si sale cara (probabilidad 0.4) gana lo que "
                   "apostó; si sale ceca, lo pierde. Termina al llegar a 100 o al quedarse sin nada.",
                   "Estado: el capital s. Acción: cuánto apostar, de 1 a min(s, 100 − s). "
                   "Recompensa: +1 al llegar a 100.",
                   "El valor de cada estado es la probabilidad de ganar."]),
            Nota("Sutton y Barto, ejemplo 4.3. A la izquierda, los valores; a la derecha, la política."),
        ], figura=apostador, ancho_figura=9.0, figura_abajo=True),

        Diapositiva("Aprender sin modelo del mundo", [
            Texto(["Casi nunca se conocen las probabilidades de transición. Hay que estimar los "
                   "valores a partir de la **experiencia**.",
                   "**Monte Carlo:** promediar los retornos observados. Hay que esperar a que termine "
                   "el episodio.",
                   "**Diferencias temporales (TD):** actualizar en cada paso, usando la estimación del "
                   "estado siguiente."]),
            Formula(r"V(S_t) \leftarrow V(S_t) + \alpha \left[ R_{t+1} + \gamma V(S_{t+1}) - V(S_t) \right]"),
            Texto(["El término entre corchetes es el **error TD**: la diferencia entre lo que se esperaba "
                   "y lo que se obtuvo más lo que se espera obtener después. α es la tasa de aprendizaje."]),
        ]),

        Diapositiva("Ejemplo: volver a casa", [
            Texto(["Al salir de la oficina se predicen 30 minutos de viaje; al ver que llueve, 40; ... "
                   "y se llega en 43. **Monte Carlo** corrige cada predicción hacia el resultado "
                   "final, y recién al llegar. **TD** corrige cada una hacia la siguiente, en el momento."]),
            Nota("Sutton y Barto, ejemplo 6.1. Si te quedás atascado en la autopista, ya sabés que "
                 "los 30 minutos eran optimistas: TD no espera a llegar para corregirlo."),
        ], figura=resumen["volver-a-casa"], ancho_figura=8.0, figura_abajo=True),

        Diapositiva("Q-learning", [
            Texto(["Aplica la idea de TD a los valores de acción, usando la ecuación de optimalidad "
                   "(Watkins, 1989):"]),
            Formula(r"Q(S_t, A_t) \leftarrow Q(S_t, A_t) + \alpha \left[ R_{t+1} + \gamma \max_a Q(S_{t+1}, a)"
                    r" - Q(S_t, A_t) \right]"),
            Texto(["No necesita modelo del mundo: aprende de las transiciones que vive.",
                   "Es **off-policy**: elige las acciones con ε-greedy, pero aprende los valores de la "
                   "política greedy, es decir, q*.",
                   "Converge a q* si todas las acciones se siguen probando en todos los estados."],
                  vinetas=True),
            Nota("Con la tabla Q aprendida, la política es elegir en cada estado la acción de mayor Q."),
        ]),

        Diapositiva("Algoritmo Q-learning", [
            Codigo("""
Parámetros: α ∈ (0, 1], ε > 0, γ ∈ [0, 1]
Inicializar Q(s, a) = 0 para todo s, a
Para cada episodio:
    Inicializar S
    Repetir para cada paso del episodio:
        Elegir A desde S con ε-greedy sobre Q
        Ejecutar A, observar R y S′
        Q(S, A) ← Q(S, A) + α [R + γ maxₐ Q(S′, a) − Q(S, A)]
        S ← S′
    hasta que S sea terminal
""", tam=12),
            Nota("Sutton y Barto, sección 6.5. Si S′ es terminal, su valor es 0."),
        ]),
    ]

    indice = [
        ("Decisiones bajo incertidumbre", 0),
        ("Utilidad y máxima utilidad esperada", 1),
        ("Exploración y explotación: tragamonedas", 1),
        ("Decisiones secuenciales", 0),
        ("Políticas y retorno", 1),
        ("Procesos de decisión de Markov", 1),
        ("Ecuaciones de Bellman", 1),
        ("Iteración de valores", 1),
        ("Aprendizaje por refuerzo sin modelo", 0),
        ("Diferencias temporales", 1),
        ("Q-learning", 1),
        ("Referencias", 0),
    ]
    referencias = [
        REF_RN, REF_SB,
        "Watkins, C. J. C. H., & Dayan, P. (1992). Q-learning. Machine Learning, 8(3–4), 279–292.",
        "Auer, P., Cesa-Bianchi, N., & Fischer, P. (2002). Finite-time analysis of the multiarmed "
        "bandit problem. Machine Learning, 47(2–3), 235–256.",
    ]
    armar("C3-1_ Decisiones secuenciales y Q-learning", "Decisiones secuenciales y Q-learning",
          indice, d, referencias)


def clase2():
    cliff = fig_cliff()
    cartpole = fig_cartpole()
    red = fig_red_q()
    memoria = fig_memoria()
    ablacion = png_de_notebook("clase2-dqn.ipynb", "variantes = {", -1, "ablacion.png")

    d = [
        Diapositiva("Repaso: Q-learning", [
            Texto(["Q-learning aprende una tabla con un valor para cada par de estado y acción:"]),
            Formula(r"Q(S_t, A_t) \leftarrow Q(S_t, A_t) + \alpha \left[ R_{t+1} + \gamma \max_a Q(S_{t+1}, a)"
                    r" - Q(S_t, A_t) \right]"),
            Texto(["Elige las acciones con ε-greedy.",
                   "Para el estado siguiente usa la **mejor** acción, aunque después no la tome.",
                   "Por eso es **off-policy**: aprende q*, el valor de la política greedy, mientras "
                   "se mueve con otra política."], vinetas=True),
        ]),

        Diapositiva("SARSA: aprender la política que se sigue", [
            Texto(["Para el estado siguiente usa la acción que el agente **efectivamente va a tomar**:"]),
            Formula(r"Q(S_t, A_t) \leftarrow Q(S_t, A_t) + \alpha \left[ R_{t+1} + \gamma\, Q(S_{t+1}, A_{t+1})"
                    r" - Q(S_t, A_t) \right]"),
            Texto(["El nombre viene de los cinco elementos de cada actualización:"]),
            Formula(r"(S_t,\ A_t,\ R_{t+1},\ S_{t+1},\ A_{t+1})"),
            Texto(["Es **on-policy**: aprende los valores de la misma política ε-greedy con la que se "
                   "mueve, incluidos los errores que comete al explorar."]),
        ]),

        Diapositiva("CliffWalking: Q-learning y SARSA", [
            Texto(["Cada paso cuesta −1 y caer al acantilado, −100. Con ε = 0.1, cada tanto el agente da "
                   "un paso al azar (Sutton y Barto, ejemplo 6.6)."]),
            Texto(["**Q-learning** aprende el camino óptimo, pegado al acantilado. Pero mientras explora "
                   "se cae seguido: recompensa promedio cercana a −50 por episodio.",
                   "**SARSA** tiene en cuenta sus propios errores y aprende un camino más alejado: "
                   "es más largo, pero mientras entrena le va mejor (cerca de −25)."], vinetas=True),
        ], figura=cliff, ancho_figura=6.6, figura_abajo=True),

        Diapositiva("Exploración decreciente", [
            Texto(["Explorar sirve al principio, cuando el agente no sabe nada, pero molesta al final. "
                   "Se empieza con un ε alto y se lo reduce con el tiempo, por ejemplo en forma lineal "
                   "durante los primeros T pasos:"]),
            Formula(r"\varepsilon_t = \max\left(\varepsilon_{\mathrm{fin}},\ \varepsilon_{\mathrm{ini}}"
                    r" - (\varepsilon_{\mathrm{ini}} - \varepsilon_{\mathrm{fin}}) \frac{t}{T}\right)"),
            Texto(["Con un ε inicial de 1 el agente arranca eligiendo todo al azar.",
                   "Termina explorando solo una fracción pequeña de las veces, por ejemplo 0.05."],
                  vinetas=True),
            Nota("Si ε llega a 0, SARSA y Q-learning terminan aprendiendo la misma política."),
        ]),

        Diapositiva("Cuando los estados no entran en una tabla", [
            Texto(["FrozenLake tiene 16 estados y CliffWalking, 48.",
                   "El backgammon tiene alrededor de 10²⁰.",
                   "En un juego de Atari el estado es la pantalla: 256^(84·84) imágenes posibles.",
                   "Si el estado son números reales (posición, velocidad), hay **infinitos** estados."],
                  vinetas=True),
            Texto(["Con una tabla aparecen dos problemas:"]),
            Texto(["**Memoria:** no entra.",
                   "**Generalización:** cada casilla se aprende por separado. Lo aprendido en un estado "
                   "no sirve para un estado parecido, y la mayoría de los estados no se visitan nunca."],
                  vinetas=True),
        ]),

        Diapositiva("Ejemplo: CartPole", [
            Texto(["Mantener un palo en equilibrio sobre un carrito."]),
            Texto(["**Estado:** posición y velocidad del carrito, ángulo y velocidad angular del palo. "
                   "Son 4 números reales.",
                   "**Acciones:** empujar a la izquierda o a la derecha.",
                   "**Recompensa:** +1 por cada paso con el palo en pie, hasta 500.",
                   "**Fin:** el palo se inclina más de 12° o el carrito se sale de la pista."],
                  vinetas=True),
            Nota("Discretizar (por ejemplo, 6 intervalos por variable: 6⁴ = 1296 estados) funciona a "
                 "medias: con pocos intervalos se pierde precisión, y con muchos cada estado se visita "
                 "muy poco."),
        ], figura=cartpole, ancho_figura=3.3),

        Diapositiva("Aproximación de funciones", [
            Texto(["Se reemplaza la tabla por una función con parámetros **w**:"]),
            Formula(r"\hat{q}(s, a; \mathbf{w}) \approx q_*(s, a)"),
            Texto(["La función tiene muchos menos parámetros que estados. Al ajustar **w** con la "
                   "experiencia de un estado, cambian también los valores de los estados parecidos: "
                   "la función **generaliza**."]),
            Nota("La tabla es un caso particular: un parámetro por cada par (s, a)."),
        ]),

        Diapositiva("Aprender como un problema de regresión", [
            Texto(["Cada transición (s, a, r, s′) da un **objetivo** para el valor de (s, a):"]),
            Formula(r"y = r + \gamma \max_{a'} \hat{q}(s', a'; \mathbf{w})"),
            Texto(["Se minimiza el error cuadrático, como en una regresión:"]),
            Formula(r"L(\mathbf{w}) = \left( y - \hat{q}(s, a; \mathbf{w}) \right)^2"),
            Texto(["con un paso de descenso por el gradiente:"]),
            Formula(r"\mathbf{w} \leftarrow \mathbf{w} + \alpha \left[ y - \hat{q}(s, a; \mathbf{w}) \right]"
                    r" \nabla_{\mathbf{w}}\, \hat{q}(s, a; \mathbf{w})"),
            Nota("y también depende de w, pero no se deriva: se lo trata como un dato. Por eso se "
                 "llama método de semi-gradiente."),
        ]),

        Diapositiva("Una red neuronal como aproximador", [
            Texto(["En DQN la función es una red neuronal:"]),
            Texto(["**Entrada:** el estado s.",
                   "**Salida:** un valor por cada acción, q̂(s, a₁), …, q̂(s, aₙ)."], vinetas=True),
            Texto(["Con una sola pasada se obtienen los valores de todas las acciones: justo lo que hace "
                   "falta para calcular el máximo del objetivo y para elegir la acción greedy."]),
            Nota("Para CartPole alcanza con dos capas ocultas de 128 neuronas: unos 17 mil parámetros."),
        ], figura=red, ancho_figura=3.3),

        Diapositiva("Q-learning con una red: por qué falla", [
            Texto(["Si se entrena la red con cada transición a medida que llega, no aprende: en "
                   "CartPole termina peor que una política al azar. Hay dos causas:"]),
            Texto(["**Muestras correlacionadas.** Las transiciones consecutivas son casi iguales. La red "
                   "se ajusta a la situación del momento y, como generaliza, cambia los valores de todos "
                   "los demás estados.",
                   "**Objetivo que se mueve.** y se calcula con la misma red que se entrena: cada paso "
                   "de gradiente cambia el objetivo."], vinetas=True),
            Nota("Ninguno de los dos problemas aparecía con la tabla."),
        ]),

        Diapositiva("La tríada mortal", [
            Texto(["Sutton y Barto llaman así a tres ingredientes que, juntos, pueden volver inestable "
                   "el aprendizaje o hacerlo diverger:"]),
            Texto(["**Aproximación de funciones.**",
                   "**Bootstrapping:** el objetivo usa estimaciones, como max q̂(s′, a′).",
                   "**Aprendizaje off-policy.**"], vinetas=True),
            Nota("Q-learning con una red tiene los tres. DQN no elimina la tríada, pero agrega dos ideas "
                 "que estabilizan mucho el entrenamiento: la memoria de repetición y la red objetivo."),
        ]),

        Diapositiva("Memoria de repetición", [
            Texto(["Cada transición (s, a, r, s′) se guarda en una memoria de capacidad fija. Cuando se "
                   "llena, se descartan las más viejas.",
                   "En cada paso se entrena con un **minibatch al azar** de la memoria."]),
            Texto(["**Rompe la correlación:** el minibatch mezcla momentos distintos.",
                   "**Reusa la experiencia:** cada transición se usa muchas veces."], vinetas=True),
            Nota("Funciona porque Q-learning es off-policy: puede aprender q* de transiciones generadas "
                 "por versiones viejas del propio agente."),
        ], figura=memoria, ancho_figura=3.3),

        Diapositiva("Red objetivo", [
            Texto(["Una copia de la red con pesos **w⁻**, congelada. Cada C pasos se copian los pesos: "
                   "w⁻ ← w. El objetivo se calcula con esa copia:"]),
            Formula(r"y = r + \gamma \max_{a'} \hat{q}(s', a'; \mathbf{w}^-)"),
            Texto(["Durante C pasos el objetivo queda fijo, y el problema se parece a una regresión común.",
                   "Juntando las dos ideas, DQN minimiza:"]),
            Formula(r"L(\mathbf{w}) = \mathbb{E}_{(s, a, r, s') \sim D} \left[ \left( r + \gamma \max_{a'}"
                    r" \hat{q}(s', a'; \mathbf{w}^-) - \hat{q}(s, a; \mathbf{w}) \right)^2 \right]"),
        ]),

        Diapositiva("Algoritmo DQN", [
            Codigo("""
Inicializar la memoria D (capacidad N), la red q̂ con pesos w al azar y w⁻ = w
Para cada paso:
    Elegir a con ε-greedy sobre q̂(s, ·; w)
    Ejecutar a, observar r y s′; guardar (s, a, r, s′, terminal) en D
    Muestrear un minibatch de D y calcular, para cada transición:
        yⱼ = rⱼ                                 si s′ⱼ es terminal
        yⱼ = rⱼ + γ maxₐ′ q̂(s′ⱼ, a′; w⁻)        si no
    Dar un paso de gradiente sobre (yⱼ − q̂(sⱼ, aⱼ; w))²
    Cada C pasos: w⁻ ← w
""", tam=11),
            Nota("Mnih et al. (2015). En la práctica se usa la pérdida de Huber y un ε decreciente, y se "
                 "empieza a entrenar cuando la memoria ya tiene algunas transiciones."),
        ]),

        Diapositiva("¿Qué aporta cada pieza?", [
            Texto(["DQN en CartPole, 100.000 pasos, tres semillas por variante:"]),
            Nota("Solo con las dos piezas juntas el agente aprende en las tres semillas. Mnih et al. "
                 "hicieron el mismo experimento con 5 juegos de Atari: allí cada idea sola ya mejoraba, "
                 "y las dos juntas mucho más."),
        ], figura=ablacion, ancho_figura=9.0, figura_abajo=True),

        Diapositiva("DQN en Atari", [
            Texto(["La entrada eran los últimos 4 cuadros de la pantalla, para percibir el movimiento.",
                   "Capas convolucionales y una salida por cada botón del joystick.",
                   "La **misma arquitectura y los mismos hiperparámetros** para 49 juegos distintos.",
                   "Unos 50 millones de cuadros por juego: alrededor de 38 días de juego.",
                   "Superó a un jugador humano profesional en 22 juegos, y alcanzó al menos el 75 % de "
                   "su puntaje en 29 de los 46 comparados."], vinetas=True),
            Nota("Sutton y Barto, sección 16.5."),
        ]),

        Diapositiva("Más allá de DQN: Double DQN", [
            Texto(["El máximo del objetivo tiende a **sobreestimar** los valores: elige la acción y la "
                   "evalúa con las mismas estimaciones ruidosas (sesgo de maximización)."]),
            Texto(["Double DQN separa las dos cosas: la red principal elige la acción y la red objetivo "
                   "la evalúa."]),
            Formula(r"y = r + \gamma\, \hat{q}\left(s', \arg\max_{a'} \hat{q}(s', a'; \mathbf{w});"
                    r" \mathbf{w}^-\right)"),
            Nota("van Hasselt, Guez y Silver (2016). Es un cambio de una línea en el código."),
        ]),
    ]

    indice = [
        ("Repaso: Q-learning y SARSA", 0),
        ("On-policy y off-policy", 1),
        ("Exploración decreciente", 1),
        ("Aproximación de funciones", 0),
        ("Cuando los estados no entran en una tabla", 1),
        ("Aprender como un problema de regresión", 1),
        ("Redes neuronales como aproximador", 1),
        ("DQN", 0),
        ("Inestabilidad y tríada mortal", 1),
        ("Memoria de repetición y red objetivo", 1),
        ("Algoritmo y resultados", 1),
        ("Referencias", 0),
    ]
    referencias = [
        REF_SB, REF_RN,
        "Lin, L.-J. (1992). Self-improving reactive agents based on reinforcement learning, planning "
        "and teaching. Machine Learning, 8(3–4), 293–321.",
        "Mnih, V., Kavukcuoglu, K., Silver, D., Graves, A., Antonoglou, I., Wierstra, D., & Riedmiller, M. "
        "(2013). Playing Atari with deep reinforcement learning. arXiv:1312.5602.",
        "Mnih, V., Kavukcuoglu, K., Silver, D., Rusu, A. A., Veness, J., Bellemare, M. G., … Hassabis, D. "
        "(2015). Human-level control through deep reinforcement learning. Nature, 518(7540), 529–533.",
        "van Hasselt, H., Guez, A., & Silver, D. (2016). Deep reinforcement learning with double "
        "Q-learning. Proceedings of the AAAI Conference on Artificial Intelligence, 30(1).",
    ]
    armar("C3-2_ Aproximación de funciones y DQN", "Aproximación de funciones y DQN",
          indice, d, referencias)


if __name__ == "__main__":
    clase1()
    clase2()
